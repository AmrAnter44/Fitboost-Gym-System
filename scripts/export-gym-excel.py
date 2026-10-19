#!/usr/bin/env python3
"""
تصدير بيانات الجيم من داتابيز SQLite إلى ملف Excel بثلاث شيتات:
  1. منتهي الاشتراك   — كل عضو تاريخ انتهاءه عدّى
  2. عليهم باقي        — كل عضو عليه مبلغ متبقي (remainingAmount > 0)
  3. الزوار            — كل الزوار غير المحذوفين + آخر متابعة

الاستخدام:
  python3 scripts/export-gym-excel.py <مسار الداتابيز> [مسار ملف الإخراج.xlsx]
  python3 scripts/export-gym-excel.py <مسار الداتابيز> [مجلد الإخراج] --split   ← كل شيت في ملف لوحده
"""
import sqlite3, sys, os
from datetime import datetime, date
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

if len(sys.argv) < 2:
    print(__doc__); sys.exit(1)

SPLIT = "--split" in sys.argv
args = [a for a in sys.argv[1:] if a != "--split"]
DB = args[0]
OUT = args[1] if len(args) > 1 else ("exports" if SPLIT else f"exports/export-{date.today().isoformat()}.xlsx")
NOW_MS = int(datetime.now().timestamp() * 1000)
TODAY = date.today()

con = sqlite3.connect(f"file:{DB}?immutable=1", uri=True)
con.row_factory = sqlite3.Row

def to_date(v):
    """Prisma/SQLite بيخزن التواريخ كـ ms أو نص ISO — نرجّع date أو None."""
    if v is None or v == "": return None
    try:
        if isinstance(v, (int, float)): return datetime.fromtimestamp(v / 1000).date()
        s = str(v).strip()
        if s.isdigit(): return datetime.fromtimestamp(int(s) / 1000).date()
        return datetime.fromisoformat(s.replace("Z", "+00:00")).date()
    except Exception:
        return None

def days_since(d):
    return (TODAY - d).days if d else None

GENDER = {"male": "ذكر", "female": "أنثى"}
VSTATUS = {"pending": "في الانتظار", "contacted": "تم التواصل", "subscribed": "اشترك",
           "not_qualified": "غير مؤهل", "rejected": "رفض"}

def member_state(r):
    if r["isBanned"]: return "محظور"
    if r["isFrozen"]: return "مجمّد"
    exp = to_date(r["expiryDate"])
    if exp and exp < TODAY: return "منتهي"
    return "نشط"

# ───────────── الاستعلامات ─────────────
MEMBER_SQL = """
SELECT m.*, c.name AS coachName, s.name AS salesName
FROM Member m
LEFT JOIN Staff c ON c.id = m.coachId
LEFT JOIN Staff s ON s.id = m.salesStaffId
"""
members = [dict(r) for r in con.execute(MEMBER_SQL)]

expired = [m for m in members if to_date(m["expiryDate"]) and to_date(m["expiryDate"]) < TODAY]
expired.sort(key=lambda m: to_date(m["expiryDate"]), reverse=True)

debtors = [m for m in members if (m["remainingAmount"] or 0) > 0]
debtors.sort(key=lambda m: -(m["remainingAmount"] or 0))

VISITOR_SQL = """
SELECT v.*,
       f.notes AS fuNotes, f.createdAt AS fuAt, f.nextFollowUpDate AS fuNext,
       f.result AS fuResult, f.salesName AS fuSales, f.stage AS fuStage,
       f.contactCount AS fuCount, st.name AS fuAssigned
FROM Visitor v
LEFT JOIN FollowUp f ON f.id = (
    SELECT id FROM FollowUp WHERE visitorId = v.id ORDER BY createdAt DESC LIMIT 1)
LEFT JOIN Staff st ON st.id = f.assignedTo
WHERE v.isDeleted = 0
ORDER BY v.createdAt DESC
"""
visitors = [dict(r) for r in con.execute(VISITOR_SQL)]

# ───────────── كتابة الإكسيل ─────────────
def new_wb():
    wb = Workbook(); wb.remove(wb.active); return wb
wb = new_wb()

HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
HEAD_FONT = Font(bold=True, color="FFFFFF", size=11)
thin = Side(style="thin", color="D9D9D9")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

SHEETS = {}
def add_sheet(title, headers, rows, widths=None, money_cols=(), date_cols=(), wb=None):
    SHEETS.setdefault(title, {"headers": headers, "rows": rows,
                              "opts": dict(widths=widths, money_cols=money_cols, date_cols=date_cols)})
    wb = wb or globals()["wb"]
    ws = wb.create_sheet(title)
    ws.sheet_view.rightToLeft = True
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font, c.border = HEAD_FILL, HEAD_FONT, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        ws.append(row)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.border = BORDER
            c.alignment = Alignment(vertical="center")
            if c.column in money_cols: c.number_format = '#,##0.00'
            if c.column in date_cols:  c.number_format = 'yyyy-mm-dd'
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, h in enumerate(headers, 1):
        w = (widths or {}).get(i)
        if not w:
            w = max([len(str(h))] + [len(str(r[i-1])) for r in rows if r[i-1] is not None][:500]) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(max(w, 8), 50)
    return ws

# 1) منتهي الاشتراك
add_sheet("منتهي الاشتراك",
    ["رقم العضو", "الاسم", "الهاتف", "هاتف احتياطي", "الجنس", "تاريخ البداية", "تاريخ الانتهاء",
     "أيام منذ الانتهاء", "سعر الاشتراك", "الباقي عليه", "الحالة", "الكوتش", "السيلز", "المصدر", "ملاحظات"],
    [[m["memberNumber"], m["name"], m["phone"], m["backupPhone"], GENDER.get(m["gender"], ""),
      to_date(m["startDate"]), to_date(m["expiryDate"]), days_since(to_date(m["expiryDate"])),
      m["subscriptionPrice"], m["remainingAmount"], member_state(m),
      m["coachName"], m["salesName"], m["source"], m["notes"]] for m in expired],
    money_cols=(9, 10), date_cols=(6, 7), widths={2: 28, 15: 40})

# 2) عليهم باقي
add_sheet("عليهم باقي",
    ["رقم العضو", "الاسم", "الهاتف", "هاتف احتياطي", "الباقي عليه", "تاريخ استحقاق الباقي",
     "سعر الاشتراك", "تاريخ الانتهاء", "حالة الاشتراك", "الكوتش", "السيلز", "ملاحظات"],
    [[m["memberNumber"], m["name"], m["phone"], m["backupPhone"], m["remainingAmount"],
      to_date(m["remainingDueDate"]), m["subscriptionPrice"], to_date(m["expiryDate"]),
      member_state(m), m["coachName"], m["salesName"], m["notes"]] for m in debtors],
    money_cols=(5, 7), date_cols=(6, 8), widths={2: 28, 12: 40})

# 3) الزوار
add_sheet("الزوار",
    ["الاسم", "الهاتف", "المصدر", "الحالة", "مهتم بـ", "الجنس", "تاريخ الزيارة",
     "آخر متابعة", "ملاحظة آخر متابعة", "نتيجة المتابعة", "المتابعة القادمة",
     "عدد مرات التواصل", "السيلز", "رقم العضو المُحيل", "ملاحظات"],
    [[v["name"], v["phone"], v["source"], VSTATUS.get(v["status"], v["status"]), v["interestedIn"],
      GENDER.get(v["gender"], ""), to_date(v["createdAt"]), to_date(v["fuAt"]), v["fuNotes"],
      v["fuResult"], to_date(v["fuNext"]), v["fuCount"], v["fuAssigned"] or v["fuSales"],
      v["referrerMemberNumber"], v["notes"]] for v in visitors],
    date_cols=(7, 8, 11), widths={1: 28, 9: 40, 15: 40})

# ملخص
def add_summary(wb):
  ws = wb.create_sheet("ملخص", 0); ws.sheet_view.rightToLeft = True
  ws.append(["البيان", "العدد"])
  for c in ws[1]: c.fill, c.font = HEAD_FILL, HEAD_FONT
  ws.append(["إجمالي الأعضاء", len(members)])
  ws.append(["منتهي الاشتراك", len(expired)])
  ws.append(["عليهم باقي", len(debtors)])
  ws.append(["إجمالي الباقي (ج.م)", round(sum(m["remainingAmount"] or 0 for m in debtors), 2)])
  ws.append(["الزوار", len(visitors)])
  ws.append(["تاريخ التصدير", datetime.now().strftime("%Y-%m-%d %H:%M")])
  ws.column_dimensions["A"].width = 24; ws.column_dimensions["B"].width = 18

if SPLIT:
    os.makedirs(OUT, exist_ok=True)
    stamp = date.today().isoformat()
    for ws in list(wb.worksheets):
        one = new_wb()
        add_sheet(ws.title, SHEETS[ws.title]["headers"], SHEETS[ws.title]["rows"],
                  wb=one, **SHEETS[ws.title]["opts"])
        path = os.path.join(OUT, f"{ws.title}-{stamp}.xlsx"); one.save(path); print("✅", path)
    one = new_wb(); add_summary(one)
    path = os.path.join(OUT, f"ملخص-{stamp}.xlsx"); one.save(path); print("✅", path)
else:
    add_summary(wb)
    os.makedirs(os.path.dirname(OUT) or ".", exist_ok=True)
    wb.save(OUT)
    print(f"✅ {OUT}")
print(f"   منتهي: {len(expired)} | عليهم باقي: {len(debtors)} | زوار: {len(visitors)}")
